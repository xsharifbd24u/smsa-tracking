import asyncio
import os
import re
from urllib.parse import urlencode
from datetime import datetime
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils.dataframe import dataframe_to_rows
import pandas as pd
from playwright.async_api import async_playwright

input_file = "tracking_numbers.xlsx"
output_file = "SMSA_Tracking_Results.xlsx"

if not os.path.exists(input_file):
    print(f"❌ Error! File '{input_file}' not found!")
    exit()

try:
    df = pd.read_excel(input_file, dtype={"Tracking Number": str})
    tracking_numbers = df["Tracking Number"].dropna().astype(str).str.strip().tolist()
    print(f"✅ Found {len(tracking_numbers)} tracking numbers\n")
except Exception as e:
    print(f"❌ Excel Reading Error: {e}")
    exit()

async def batch_track():
    print("🔄 Starting tracking...\n")
    start_time = datetime.now()
    print(f"⏰ Started: {start_time.strftime('%Y-%m-%d %H:%M:%S')}\n")
    
    batch_size = 20
    all_results = []
    
    for batch_start in range(0, len(tracking_numbers), batch_size):
        batch = tracking_numbers[batch_start:batch_start + batch_size]
        print(f"\n📦 Batch {batch_start//batch_size + 1}: Processing {len(batch)} numbers...")
        
        query_params = {f"tracknumbers[{i}]": num for i, num in enumerate(batch)}
        target_url = "https://www.smsaexpress.com/trackingdetails?" + urlencode(query_params)
        
        async with async_playwright() as p:
            browser = await p.chromium.launch(headless=True)
            context = await browser.new_context(
                user_agent="Mozilla/5.0 (Windows NT 10.0; Win64; x64) AppleWebKit/537.36 (KHTML, like Gecko) Chrome/122.0.0.0 Safari/537.36"
            )
            page = await context.new_page()
            
            try:
                print(f"🌐 Loading SMSA website (Batch {batch_start//batch_size + 1})...")
                await page.goto(target_url, wait_until="networkidle", timeout=60000)
                await page.wait_for_timeout(3000)
                
                page_text = await page.inner_text("body")
                
                for track_id in batch:
                    print(f"   📦 {track_id}")
                    
                    if track_id in page_text:
                        idx = page_text.find(track_id)
                        block = page_text[idx:idx + 1200]
                        
                        # ---- Find Status ----
                        status = "In Transit"
                        status_match = re.search(r"Status\s*\n\s*([^\n]+(?:Delivered|Out for Delivery|In Transit|Shipment Received)[^\n]*)", block, re.IGNORECASE)
                        if status_match:
                            status = status_match.group(1).strip()
                        else:
                            delivered_match = re.search(r"Delivered,?\s*Signed by:\s*([^\n]+)", block, re.IGNORECASE)
                            if delivered_match:
                                status = f"Delivered, Signed by: {delivered_match.group(1).strip()}"
                            elif "Out for Delivery" in block:
                                status = "Out for Delivery"
                            elif "Shipment Received at SMSA Sorting Facility" in block:
                                status = "Shipment Received at SMSA Sorting Facility"
                            elif "In Transit" in block:
                                status = "In Transit"
                            else:
                                status_match2 = re.search(r"(?:Status|STATUS)\s*\n\s*([^\n]+)", block, re.IGNORECASE)
                                if status_match2:
                                    status = status_match2.group(1).strip()
                        
                        # ---- Determine Status Category ----
                        if "Delivered" in status or "delivered" in status:
                            status_category = "Delivered"
                        elif "Out for Delivery" in status or "Shipment Received" in status or "In Transit" in status or "Departed" in status or "Arrived" in status:
                            status_category = "In Transit"
                        elif "Picked Up" in status or "Picked up" in status:
                            status_category = "Picked Up"
                        else:
                            status_category = "In Transit"
                        
                        # ---- Find From ----
                        origin = "N/A"
                        from_match = re.search(r"From\s*\n\s*([^\n]+)", block, re.IGNORECASE)
                        if from_match:
                            origin = from_match.group(1).strip()
                            origin = re.sub(r'^From\s*', '', origin)
                            if "mé & Príncipe" in origin or "Sao Tome" in origin:
                                origin = "N/A"
                        
                        # ---- Find To ----
                        destination = "N/A"
                        to_match = re.search(r"To\s*\n\s*([^\n]+)", block, re.IGNORECASE)
                        if to_match:
                            destination = to_match.group(1).strip()
                            destination = re.sub(r'^To\s*', '', destination)
                            if "mé & Príncipe" in destination or "Sao Tome" in destination:
                                destination = "N/A"
                        
                        if origin == "N/A":
                            from_alt = re.search(r"From\s+([^\n]+)", block, re.IGNORECASE)
                            if from_alt:
                                origin = from_alt.group(1).strip()
                                if "mé & Príncipe" in origin:
                                    origin = "N/A"
                        
                        if destination == "N/A":
                            to_alt = re.search(r"To\s+([^\n]+)", block, re.IGNORECASE)
                            if to_alt:
                                destination = to_alt.group(1).strip()
                                if "mé & Príncipe" in destination:
                                    destination = "N/A"
                        
                        # ---- Find Last Update ----
                        last_update = "N/A"
                        date_match = re.search(r"(\d{2}-[A-Za-z]{3}-\d{4})\s*\n\s*([A-Za-z]+)?", block, re.IGNORECASE)
                        if date_match:
                            last_update = date_match.group(1).strip()
                            if date_match.group(2):
                                last_update += f" ({date_match.group(2).strip()})"
                        
                        time_match = re.search(r"(\d{2}:\d{2}\s*[AP]M)", block, re.IGNORECASE)
                        if time_match and last_update != "N/A":
                            last_update += f" {time_match.group(1).strip()}"
                        
                        # ---- Check if shipment has moved from pickup (has updates) ----
                        # If there's any tracking progress beyond pickup, it's updated
                        has_updates = False
                        if "Out for Delivery" in block or "Shipment Received" in block or "Departed" in block or "Arrived" in block or "Delivered" in block:
                            has_updates = True
                        
                        all_results.append({
                            "Tracking Number": track_id,
                            "Status": status,
                            "Status Category": status_category,
                            "From": origin,
                            "To": destination,
                            "Last Update": last_update,
                            "Has Updates": has_updates,
                            "Check Result": "✅ Success"
                        })
                        print(f"      ✅ Status: {status[:50]}...")
                        print(f"      📍 From: {origin} → To: {destination}")
                        
                    else:
                        all_results.append({
                            "Tracking Number": track_id,
                            "Status": "❌ Not Found",
                            "Status Category": "Not Found",
                            "From": "N/A",
                            "To": "N/A",
                            "Last Update": "N/A",
                            "Has Updates": False,
                            "Check Result": "❌ Failed"
                        })
                        print(f"      ❌ Not Found")
                        
            except Exception as e:
                print(f"   ❌ Batch Processing Error: {e}")
                for track_id in batch:
                    all_results.append({
                        "Tracking Number": track_id,
                        "Status": "⚠️ Error",
                        "Status Category": "Error",
                        "From": "N/A",
                        "To": "N/A",
                        "Last Update": "N/A",
                        "Has Updates": False,
                        "Check Result": "❌ Failed"
                    })
            
            await browser.close()
    
    end_time = datetime.now()
    duration = end_time - start_time
    print(f"\n⏰ Finished: {end_time.strftime('%Y-%m-%d %H:%M:%S')}")
    print(f"⏱️ Total Time: {duration.total_seconds():.2f} seconds")
    
    create_excel(all_results, start_time, end_time)

def create_excel(results, start_time, end_time):
    try:
        print("\n📊 Creating Excel file with color coding...")
        result_df = pd.DataFrame(results)
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "SMSA Tracking Status"
        
        # Title
        ws.merge_cells('A1:H1')
        title_cell = ws['A1']
        title_cell.value = "📦 SMSA EXPRESS TRACKING REPORT"
        title_cell.font = Font(name="Calibri", size=16, bold=True, color="1F497D")
        title_cell.alignment = Alignment(horizontal="center", vertical="center")
        
        # Info with color legend
        ws.merge_cells('A2:H2')
        info_cell = ws['A2']
        info_cell.value = f"Generated: {start_time.strftime('%Y-%m-%d %H:%M:%S')} | Duration: {(end_time - start_time).total_seconds():.2f}s | Total Records: {len(results)}"
        info_cell.font = Font(name="Calibri", size=10, color="666666")
        info_cell.alignment = Alignment(horizontal="center", vertical="center")
        
        # Color Legend (Row 3)
        ws.merge_cells('A3:H3')
        legend_cell = ws['A3']
        legend_cell.value = "🟢 Delivered  |  🟡 In Transit / Out for Delivery / Updated  |  🔴 Picked Up Only (Not Moved)"
        legend_cell.font = Font(name="Calibri", size=10, bold=True)
        legend_cell.alignment = Alignment(horizontal="center", vertical="center")
        
        # Headers
        headers = ["Tracking Number", "Tracking Link", "Status", "Status Category", "From", "To", "Last Update", "Check Result"]
        for col_idx, header in enumerate(headers, 1):
            cell = ws.cell(row=4, column=col_idx)
            cell.value = header
        
        # Color Definitions
        green_fill = PatternFill(start_color="C6EFCE", end_color="C6EFCE", fill_type="solid")  # Light Green
        green_font = Font(color="006100", bold=True)
        
        yellow_fill = PatternFill(start_color="FFEB9C", end_color="FFEB9C", fill_type="solid")  # Light Yellow
        yellow_font = Font(color="9C5700", bold=True)
        
        red_fill = PatternFill(start_color="FFC7CE", end_color="FFC7CE", fill_type="solid")  # Light Red
        red_font = Font(color="9C0006", bold=True)
        
        # Styles
        header_fill = PatternFill(start_color="1F497D", end_color="1F497D", fill_type="solid")
        header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
        row_alt_fill = PatternFill(start_color="F2F5F9", end_color="F2F5F9", fill_type="solid")
        thin_border = Border(
            left=Side(style="thin", color="D9D9D9"),
            right=Side(style="thin", color="D9D9D9"),
            top=Side(style="thin", color="D9D9D9"),
            bottom=Side(style="thin", color="D9D9D9"),
        )
        
        # Insert data with conditional coloring
        for row_idx, record in enumerate(results, start=5):
            # Determine color based on status category
            if record["Status Category"] == "Delivered":
                row_fill = green_fill
                row_font = green_font
            elif record["Status Category"] == "In Transit" or record["Has Updates"] == True:
                row_fill = yellow_fill
                row_font = yellow_font
            elif record["Status Category"] == "Picked Up":
                row_fill = red_fill
                row_font = red_font
            else:
                row_fill = None
                row_font = None
            
            # Tracking Number
            num_cell = ws.cell(row=row_idx, column=1, value=record["Tracking Number"])
            if row_fill:
                num_cell.fill = row_fill
            if row_font:
                num_cell.font = row_font
            
            # Tracking Link (Hyperlink)
            link_url = f"https://www.smsaexpress.com/trackingdetails?tracknumbers[0]={record['Tracking Number']}"
            link_cell = ws.cell(row=row_idx, column=2, value="🔗 View")
            link_cell.hyperlink = link_url
            link_cell.font = Font(color="0563C1", underline="single")
            if row_fill:
                link_cell.fill = row_fill
            link_cell.alignment = Alignment(horizontal="center", vertical="center")
            
            # Status
            status_cell = ws.cell(row=row_idx, column=3, value=record["Status"])
            if row_fill:
                status_cell.fill = row_fill
            if row_font:
                status_cell.font = row_font
            
            # Status Category
            cat_cell = ws.cell(row=row_idx, column=4, value=record["Status Category"])
            if row_fill:
                cat_cell.fill = row_fill
            if row_font:
                cat_cell.font = row_font
            
            # From
            from_cell = ws.cell(row=row_idx, column=5, value=record["From"])
            if row_fill:
                from_cell.fill = row_fill
            if row_font:
                from_cell.font = row_font
            
            # To
            to_cell = ws.cell(row=row_idx, column=6, value=record["To"])
            if row_fill:
                to_cell.fill = row_fill
            if row_font:
                to_cell.font = row_font
            
            # Last Update
            update_cell = ws.cell(row=row_idx, column=7, value=record["Last Update"])
            if row_fill:
                update_cell.fill = row_fill
            if row_font:
                update_cell.font = row_font
            
            # Check Result
            result_cell = ws.cell(row=row_idx, column=8, value=record["Check Result"])
            if row_fill:
                result_cell.fill = row_fill
            if row_font:
                result_cell.font = row_font
        
        # Header styles
        for col_idx in range(1, 9):
            cell = ws.cell(row=4, column=col_idx)
            cell.fill = header_fill
            cell.font = header_font
            cell.alignment = Alignment(horizontal="center", vertical="center")
            cell.border = thin_border
        
        # Data row styles (borders and alignment)
        for row_idx in range(5, ws.max_row + 1):
            for col_idx in range(1, 9):
                cell = ws.cell(row=row_idx, column=col_idx)
                cell.border = thin_border
                if col_idx in [1, 2, 4, 7, 8]:
                    cell.alignment = Alignment(horizontal="center", vertical="center")
                else:
                    cell.alignment = Alignment(horizontal="left", vertical="center")
        
        # Column widths
        column_widths = {'A': 18, 'B': 15, 'C': 45, 'D': 20, 'E': 25, 'F': 25, 'G': 20, 'H': 15}
        for col_letter, width in column_widths.items():
            ws.column_dimensions[col_letter].width = width
        
        # Footer
        footer_row = ws.max_row + 1
        ws.merge_cells(f'A{footer_row}:H{footer_row}')
        footer_cell = ws[f'A{footer_row}']
        footer_cell.value = f"🟢 Delivered | 🟡 In Transit/Updated | 🔴 Picked Up Only | Report: {datetime.now().strftime('%Y-%m-%d %H:%M:%S')}"
        footer_cell.font = Font(name="Calibri", size=9, color="666666")
        footer_cell.alignment = Alignment(horizontal="center", vertical="center")
        
        wb.save(output_file)
        print(f"\n✅ Success! File saved: '{output_file}'")
        print(f"📁 Location: {os.path.abspath(output_file)}")
        print(f"📊 Total {len(results)} records processed")
        
        # Summary with color counts
        delivered = sum(1 for r in results if r["Status Category"] == "Delivered")
        in_transit = sum(1 for r in results if r["Status Category"] == "In Transit" or r.get("Has Updates", False))
        picked_up = sum(1 for r in results if r["Status Category"] == "Picked Up")
        failed = sum(1 for r in results if r["Check Result"] != "✅ Success")
        
        print(f"\n📊 Color Summary:")
        print(f"   🟢 Delivered: {delivered}")
        print(f"   🟡 In Transit/Updated: {in_transit}")
        print(f"   🔴 Picked Up Only: {picked_up}")
        print(f"   ❌ Failed: {failed}")
        print(f"   ⏱️ Time: {(end_time - start_time).total_seconds():.2f} seconds")
        
    except PermissionError:
        print(f"\n❌ Error! File '{output_file}' is open in Excel.")
        print("👉 Please close the file and try again.")

if __name__ == "__main__":
    asyncio.run(batch_track())
